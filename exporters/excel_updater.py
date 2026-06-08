"""
Atualiza uma planilha Excel fixa com os dados mais recentes do banco.

Diferente do ExcelExporter (que gera arquivos com timestamp), este módulo
mantém um arquivo único que é sobrescrito a cada execução — ideal para
conectar ao Power BI Desktop via "Atualizar".

Estrutura da planilha:
    Aba 'Empenhos'          → fato_empenho completo
    Aba 'Fichas'            → fato_ficha_financeira completo
    Aba 'Consolidado'       → execução por UG/Ação/Subação (view Power BI)
    Aba 'Top Credores'      → 20 maiores credores por valor empenhado
    Aba 'Por Natureza'      → fichas agrupadas por natureza de despesa
    Aba 'Resumo'            → totalizadores gerais + data de atualização
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from utils.helpers import garantir_diretorio
from utils.logger import setup_logger

logger = setup_logger(__name__)

# Paleta visual
COR_AZUL_ESCURO  = "1F4E79"
COR_AZUL_MEDIO   = "2E75B6"
COR_AZUL_CLARO   = "BDD7EE"
COR_VERDE        = "375623"
COR_VERDE_CLARO  = "E2EFDA"
COR_AMARELO      = "FFC000"
COR_CINZA_CLARO  = "F2F2F2"
COR_BRANCO       = "FFFFFF"
COR_TEXTO_CLARO  = "FFFFFF"
COR_BORDA        = "B8CCE4"


class ExcelUpdater:
    """
    Mantém uma planilha Excel fixa atualizada com os dados do e-Fisco.

    Args:
        caminho: Path completo do arquivo .xlsx a criar ou sobrescrever.
    """

    def __init__(self, caminho: Path) -> None:
        self._caminho = caminho
        garantir_diretorio(caminho.parent)

    # ------------------------------------------------------------------
    # Interface pública
    # ------------------------------------------------------------------

    def atualizar(
        self,
        df_empenhos: pd.DataFrame,
        df_fichas: pd.DataFrame,
    ) -> Path:
        """
        Recria a planilha com os DataFrames fornecidos.

        Args:
            df_empenhos: Dados de fato_empenho.
            df_fichas:   Dados de fato_ficha_financeira.

        Returns:
            Path do arquivo salvo.
        """
        logger.info("Gerando planilha Excel: %s", self._caminho)

        df_consolidado   = self._consolidado(df_empenhos)
        df_top_credores  = self._top_credores(df_empenhos)
        df_por_natureza  = self._por_natureza(df_fichas)

        wb = Workbook()
        wb.remove(wb.active)

        self._aba_resumo(wb, df_empenhos, df_fichas, df_consolidado)
        self._aba_dados(wb, "Empenhos",    df_empenhos,    COR_AZUL_ESCURO)
        self._aba_dados(wb, "Fichas",      df_fichas,      COR_VERDE)
        self._aba_dados(wb, "Consolidado", df_consolidado, COR_AZUL_MEDIO)
        self._aba_dados(wb, "Top Credores", df_top_credores, COR_AZUL_MEDIO)
        self._aba_dados(wb, "Por Natureza", df_por_natureza, COR_VERDE)

        wb.save(self._caminho)
        logger.info(
            "Planilha salva: %s  |  empenhos=%d  fichas=%d",
            self._caminho,
            len(df_empenhos),
            len(df_fichas),
        )
        return self._caminho

    # ------------------------------------------------------------------
    # Transformações analíticas
    # ------------------------------------------------------------------

    @staticmethod
    def _consolidado(df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            return pd.DataFrame(columns=[
                "unidade_gestora", "acao", "subacao", "data_coleta",
                "total_empenhado", "total_liquidado", "total_pago",
                "saldo", "pct_execucao",
            ])
        grp = df.groupby(
            ["unidade_gestora", "acao", "subacao", "data_coleta"], as_index=False
        ).agg(
            total_empenhado=("valor_empenhado", "sum"),
            total_liquidado=("valor_liquidado", "sum"),
            total_pago=("valor_pago", "sum"),
        )
        grp["saldo"] = grp["total_empenhado"] - grp["total_pago"]
        grp["pct_execucao"] = grp.apply(
            lambda r: round(r["total_pago"] / r["total_empenhado"] * 100, 2)
            if r["total_empenhado"] > 0 else 0.0,
            axis=1,
        )
        return grp.sort_values("total_empenhado", ascending=False)

    @staticmethod
    def _top_credores(df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            return pd.DataFrame(columns=[
                "credor", "qtd_empenhos", "total_empenhado", "total_pago",
            ])
        return (
            df.groupby("credor", as_index=False)
            .agg(
                qtd_empenhos=("credor", "count"),
                total_empenhado=("valor_empenhado", "sum"),
                total_liquidado=("valor_liquidado", "sum"),
                total_pago=("valor_pago", "sum"),
            )
            .sort_values("total_empenhado", ascending=False)
            .head(20)
        )

    @staticmethod
    def _por_natureza(df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            return pd.DataFrame(columns=[
                "natureza_despesa", "grupo_despesa", "fonte_recurso",
                "dotacao_atual", "valor_empenhado", "valor_liquidado",
                "valor_pago", "saldo",
            ])
        return (
            df.groupby(
                ["natureza_despesa", "grupo_despesa", "fonte_recurso"],
                as_index=False,
                dropna=False,
            )
            .agg(
                dotacao_inicial=("dotacao_inicial", "sum"),
                dotacao_atual=("dotacao_atual", "sum"),
                valor_empenhado=("valor_empenhado", "sum"),
                valor_liquidado=("valor_liquidado", "sum"),
                valor_pago=("valor_pago", "sum"),
                saldo=("saldo", "sum"),
            )
            .sort_values("valor_empenhado", ascending=False)
        )

    # ------------------------------------------------------------------
    # Construção das abas
    # ------------------------------------------------------------------

    def _aba_resumo(
        self,
        wb: Workbook,
        df_emp: pd.DataFrame,
        df_fic: pd.DataFrame,
        df_cons: pd.DataFrame,
    ) -> None:
        ws = wb.create_sheet("Resumo", 0)
        ws.sheet_view.showGridLines = False

        # Cabeçalho principal
        self._mesclar_titulo(ws, "EXECUÇÃO ORÇAMENTÁRIA – e-Fisco PE", "A1", "F1", COR_AZUL_ESCURO)
        self._mesclar_titulo(ws, f"Atualizado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}", "A2", "F2", COR_AZUL_MEDIO, tamanho=11)

        # Totalizadores
        linha = 4
        total_emp = df_emp["valor_empenhado"].sum() if not df_emp.empty else 0
        total_liq = df_emp["valor_liquidado"].sum() if not df_emp.empty else 0
        total_pago = df_emp["valor_pago"].sum() if not df_emp.empty else 0
        saldo = total_emp - total_pago
        pct = round(total_pago / total_emp * 100, 2) if total_emp > 0 else 0

        cards = [
            ("Total Empenhado",  total_emp,  COR_AZUL_ESCURO),
            ("Total Liquidado",  total_liq,  COR_AZUL_MEDIO),
            ("Total Pago",       total_pago, COR_VERDE),
            ("Saldo a Pagar",    saldo,      COR_AMARELO),
            ("% Execução",       f"{pct}%",  COR_AZUL_ESCURO),
        ]

        colunas = ["A", "B", "C", "D", "E"]
        for col, (titulo, valor, cor) in zip(colunas, cards):
            cell_t = ws[f"{col}{linha}"]
            cell_t.value = titulo
            cell_t.fill = PatternFill("solid", fgColor=cor)
            cell_t.font = Font(color=COR_BRANCO, bold=True, size=10)
            cell_t.alignment = Alignment(horizontal="center", vertical="center")

            cell_v = ws[f"{col}{linha+1}"]
            cell_v.value = (
                valor if isinstance(valor, str)
                else f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
            )
            cell_v.fill = PatternFill("solid", fgColor=COR_AZUL_CLARO)
            cell_v.font = Font(bold=True, size=12, color=COR_AZUL_ESCURO)
            cell_v.alignment = Alignment(horizontal="center", vertical="center")

            ws.row_dimensions[linha].height = 20
            ws.row_dimensions[linha + 1].height = 28
            ws.column_dimensions[col].width = 22

        # Mini-tabela consolidada (top 10 linhas)
        linha_tab = linha + 4
        self._mesclar_titulo(ws, "CONSOLIDADO POR AÇÃO/SUBAÇÃO (Top 10)", f"A{linha_tab}", f"F{linha_tab}", COR_AZUL_ESCURO, tamanho=11)
        linha_tab += 1

        cabecalhos_cons = [
            "UG", "Ação", "Subação", "Empenhado (R$)", "Pago (R$)", "% Exec."
        ]
        for j, cab in enumerate(cabecalhos_cons, 1):
            cell = ws.cell(row=linha_tab, column=j, value=cab)
            cell.fill = PatternFill("solid", fgColor=COR_AZUL_MEDIO)
            cell.font = Font(color=COR_BRANCO, bold=True, size=10)
            cell.alignment = Alignment(horizontal="center")

        if not df_cons.empty:
            for i, row in enumerate(df_cons.head(10).itertuples(index=False), 1):
                cor_linha = COR_AZUL_CLARO if i % 2 == 0 else COR_BRANCO
                valores = [
                    row.unidade_gestora, row.acao, row.subacao,
                    row.total_empenhado, row.total_pago, row.pct_execucao,
                ]
                for j, val in enumerate(valores, 1):
                    cell = ws.cell(row=linha_tab + i, column=j, value=val)
                    cell.fill = PatternFill("solid", fgColor=cor_linha)
                    cell.alignment = Alignment(horizontal="center")
                    if j in (4, 5):
                        cell.number_format = '#,##0.00'

        ws.freeze_panes = "A3"

    def _aba_dados(
        self,
        wb: Workbook,
        nome: str,
        df: pd.DataFrame,
        cor_cabecalho: str,
    ) -> None:
        ws = wb.create_sheet(nome)
        ws.sheet_view.showGridLines = False

        if df.empty:
            ws["A1"] = "(sem dados para esta aba)"
            ws["A1"].font = Font(italic=True, color="808080")
            return

        # Cabeçalho
        for col_idx, coluna in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=str(coluna).upper().replace("_", " "))
            cell.fill = PatternFill("solid", fgColor=cor_cabecalho)
            cell.font = Font(color=COR_BRANCO, bold=True, size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            borda = Border(
                bottom=Side(border_style="medium", color=COR_BRANCO),
                right=Side(border_style="thin", color=COR_BRANCO),
            )
            cell.border = borda

        ws.row_dimensions[1].height = 30

        # Dados com zebra
        colunas_monetarias = {
            col for col in df.columns
            if any(x in col for x in ("valor", "dotacao", "saldo", "empenhado", "liquidado", "pago", "total"))
        }
        colunas_pct = {col for col in df.columns if "pct" in col}

        for row_idx, row_data in enumerate(df.itertuples(index=False, name=None), 2):
            cor_fundo = COR_CINZA_CLARO if row_idx % 2 == 0 else COR_BRANCO
            for col_idx, valor in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=valor)
                cell.fill = PatternFill("solid", fgColor=cor_fundo)
                cell.alignment = Alignment(vertical="center")

                col_nome = df.columns[col_idx - 1]
                if col_nome in colunas_monetarias:
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_nome in colunas_pct:
                    cell.number_format = '0.00"%"'
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        # Linha de totais para colunas numéricas
        linha_total = len(df) + 2
        cell_label = ws.cell(row=linha_total, column=1, value="TOTAL")
        cell_label.font = Font(bold=True, color=COR_BRANCO)
        cell_label.fill = PatternFill("solid", fgColor=COR_AZUL_ESCURO)
        cell_label.alignment = Alignment(horizontal="center")

        for col_idx, col_nome in enumerate(df.columns, 1):
            if col_nome in colunas_monetarias:
                cell_tot = ws.cell(row=linha_total, column=col_idx)
                col_letra = get_column_letter(col_idx)
                cell_tot.value = f"=SUM({col_letra}2:{col_letra}{linha_total-1})"
                cell_tot.number_format = '#,##0.00'
                cell_tot.font = Font(bold=True, color=COR_BRANCO)
                cell_tot.fill = PatternFill("solid", fgColor=COR_AZUL_ESCURO)
                cell_tot.alignment = Alignment(horizontal="right")
            else:
                cell = ws.cell(row=linha_total, column=col_idx)
                cell.fill = PatternFill("solid", fgColor=COR_AZUL_ESCURO)

        # Largura das colunas
        for col_idx, col_nome in enumerate(df.columns, 1):
            max_len = max(
                len(str(col_nome)),
                df[col_nome].astype(str).str.len().max() if not df.empty else 0,
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 45)

        # Filtro automático no cabeçalho
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}1"
        ws.freeze_panes = "A2"

    @staticmethod
    def _mesclar_titulo(
        ws,
        texto: str,
        cel_ini: str,
        cel_fim: str,
        cor: str,
        tamanho: int = 14,
    ) -> None:
        ws.merge_cells(f"{cel_ini}:{cel_fim}")
        cell = ws[cel_ini]
        cell.value = texto
        cell.fill = PatternFill("solid", fgColor=cor)
        cell.font = Font(color=COR_BRANCO, bold=True, size=tamanho)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        linha = int("".join(filter(str.isdigit, cel_ini)))
        ws.row_dimensions[linha].height = 28
