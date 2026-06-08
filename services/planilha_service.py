"""
Serviço de exportação para a planilha operacional existente.

Estratégia:
  1. Copia o arquivo-base para o destino configurado (preserva estrutura/formato)
  2. Atualiza a aba SITUAÇÃO_DOS_DEA'S com os empenhos frescos do e-Fisco
  3. Adiciona/substitui as abas eFISCO_Empenhos, eFISCO_Fichas e eFISCO_Consolidado
     com os dados brutos para uso no Power BI ou análise ad-hoc
"""
from __future__ import annotations

import shutil
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from config.settings import settings
from database.connection import obter_sessao
from database.models import FatoEmpenho, FatoFichaFinanceira
from utils.helpers import garantir_diretorio
from utils.logger import setup_logger

logger = setup_logger(__name__)

# Caminho da planilha-base (template original, nunca modificado)
PLANILHA_BASE_DEFAULT = Path(__file__).parent.parent / "base" / "planilha_base.xlsx"

# Paleta visual
_AZ_ESCURO = "1F4E79"
_AZ_MEDIO  = "2E75B6"
_AZ_CLARO  = "BDD7EE"
_VERDE     = "375623"
_VERDE_CL  = "E2EFDA"
_CINZA     = "F2F2F2"
_BRANCO    = "FFFFFF"


class PlanilhaService:
    """
    Atualiza a planilha operacional com os dados mais recentes do e-Fisco.

    Args:
        planilha_base:   Path do arquivo-base (template). Nunca é alterado.
        planilha_saida:  Path do arquivo de destino (será sobrescrito).
    """

    def __init__(
        self,
        planilha_base: Path,
        planilha_saida: Path | None = None,
    ) -> None:
        self._base = planilha_base
        self._saida = planilha_saida or (
            garantir_diretorio(settings.export.diretorio) / "efisco_dados.xlsx"
        )

    # ------------------------------------------------------------------
    # Interface pública
    # ------------------------------------------------------------------

    def exportar(self, data_coleta: date | None = None) -> Path:
        """
        Gera a planilha de saída com dados do banco.

        Args:
            data_coleta: Data de referência (padrão: hoje).

        Returns:
            Path do arquivo gerado.
        """
        data_coleta = data_coleta or date.today()
        logger.info("Iniciando exportação para planilha: %s", self._saida)

        if not self._base.exists():
            raise FileNotFoundError(
                f"Planilha-base não encontrada: {self._base}\n"
                "Coloque o arquivo original em 'base/planilha_base.xlsx' "
                "ou passe o caminho via --planilha-base."
            )

        # Copia o template preservando toda a estrutura e formatação
        garantir_diretorio(self._saida.parent)
        shutil.copy2(self._base, self._saida)
        logger.info("Template copiado para: %s", self._saida)

        df_empenhos = self._carregar_empenhos(data_coleta)
        df_fichas   = self._carregar_fichas(data_coleta)

        wb = load_workbook(self._saida)

        self._atualizar_deas(wb, df_empenhos, data_coleta)
        self._injetar_aba_efisco(wb, "eFISCO_Empenhos",    df_empenhos,   _AZ_ESCURO)
        self._injetar_aba_efisco(wb, "eFISCO_Fichas",      df_fichas,     _VERDE)
        self._injetar_aba_efisco(wb, "eFISCO_Consolidado", self._consolidado(df_empenhos), _AZ_MEDIO)

        wb.save(self._saida)
        logger.info(
            "Planilha salva: %s | empenhos=%d fichas=%d",
            self._saida, len(df_empenhos), len(df_fichas),
        )
        return self._saida

    # ------------------------------------------------------------------
    # Carga de dados do banco
    # ------------------------------------------------------------------

    @staticmethod
    def _carregar_empenhos(data_coleta: date) -> pd.DataFrame:
        try:
            from sqlalchemy import select
            with obter_sessao() as session:
                rows = session.scalars(
                    select(FatoEmpenho).where(FatoEmpenho.data_coleta == data_coleta)
                ).all()
            if not rows:
                logger.warning("Sem empenhos para %s. DataFrame vazio.", data_coleta)
                return pd.DataFrame()
            return pd.DataFrame([{
                "data_coleta":          r.data_coleta,
                "unidade_gestora":      r.unidade_gestora,
                "acao":                 r.acao,
                "subacao":              r.subacao,
                "data_empenho":         r.data_empenho,
                "credor":               r.credor,
                "valor_empenhado_item": float(r.valor_empenhado_item),
                "valor_empenhado":      float(r.valor_empenhado),
                "valor_liquidado":      float(r.valor_liquidado),
                "valor_pago":           float(r.valor_pago),
            } for r in rows])
        except Exception as exc:
            logger.error("Erro ao carregar empenhos: %s — usando DataFrame vazio.", exc)
            return pd.DataFrame()

    @staticmethod
    def _carregar_fichas(data_coleta: date) -> pd.DataFrame:
        try:
            from sqlalchemy import select
            with obter_sessao() as session:
                rows = session.scalars(
                    select(FatoFichaFinanceira).where(
                        FatoFichaFinanceira.data_coleta == data_coleta
                    )
                ).all()
            if not rows:
                logger.warning("Sem fichas para %s. DataFrame vazio.", data_coleta)
                return pd.DataFrame()
            return pd.DataFrame([{
                "data_coleta":          r.data_coleta,
                "ficha_id":             r.ficha_id,
                "exercicio":            r.exercicio,
                "unidade_gestora":      r.unidade_gestora,
                "gestao":               r.gestao,
                "grupo_despesa":        r.grupo_despesa,
                "fonte_recurso":        r.fonte_recurso,
                "natureza_despesa":     r.natureza_despesa,
                "dea":                  r.dea,
                "detalhamento_despesa": r.detalhamento_despesa,
                "destinacao_recurso":   r.destinacao_recurso,
                "situacao":             r.situacao,
                "dotacao_inicial":      float(r.dotacao_inicial),
                "dotacao_atual":        float(r.dotacao_atual),
                "valor_empenhado":      float(r.valor_empenhado),
                "valor_liquidado":      float(r.valor_liquidado),
                "valor_pago":           float(r.valor_pago),
                "saldo":                float(r.saldo),
            } for r in rows])
        except Exception as exc:
            logger.error("Erro ao carregar fichas: %s — usando DataFrame vazio.", exc)
            return pd.DataFrame()

    # ------------------------------------------------------------------
    # Atualização da aba SITUAÇÃO_DOS_DEA'S
    # ------------------------------------------------------------------

    def _atualizar_deas(
        self,
        wb,
        df: pd.DataFrame,
        data_coleta: date,
    ) -> None:
        """
        Substitui os dados de DEA preservando título e cabeçalho.

        Estrutura original:
          L1 – título com data
          L2 – cabeçalho (EMPENHO, DATA NE, CREDOR, VALOR EMPENHADO, ...)
          L3+ – dados
          Última – totais
        """
        # Nome da aba (tem caractere especial ´)
        nome_aba = next(
            (n for n in wb.sheetnames if "DEA" in n.upper()), None
        )
        if not nome_aba:
            logger.warning("Aba SITUAÇÃO_DOS_DEA'S não encontrada; pulando.")
            return

        ws = wb[nome_aba]

        # Atualiza o título com a data atual
        data_fmt = data_coleta.strftime("%d/%m/%Y")
        celula_titulo = ws["B1"]
        celula_titulo.value = f"SITUAÇÃO DOS DEA'S EM {data_fmt} CONFORME O E-FISCO"

        # Limpa linhas de dados (L3 até o fim) — ignora células mescladas
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
            for cell in row:
                try:
                    cell.value = None
                except AttributeError:
                    pass  # MergedCell — ignorar

        if df.empty:
            logger.info("Sem empenhos para atualizar DEA.")
            return

        # Preenche dados
        fill_par  = PatternFill("solid", fgColor="EBF3FB")
        fill_impar = PatternFill("solid", fgColor=_BRANCO)
        fonte_dados = Font(size=10)
        borda = Border(
            bottom=Side(border_style="thin", color="B8CCE4"),
            right=Side(border_style="thin", color="B8CCE4"),
        )

        for i, row in enumerate(df.itertuples(index=False), start=3):
            fill = fill_par if i % 2 == 0 else fill_impar
            dados_linha = [
                None,                         # col A vazia (segue padrão original)
                f"{row.data_empenho.year if row.data_empenho else ''}NE",  # NE placeholder
                row.data_empenho,
                row.credor,
                row.valor_empenhado,
                row.valor_liquidado,
                row.valor_pago,
            ]
            for j, val in enumerate(dados_linha, 1):
                cell = ws.cell(row=i, column=j, value=val)
                cell.fill = fill
                cell.font = fonte_dados
                cell.border = borda
                if j in (5, 6, 7):
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif j == 3 and isinstance(val, date):
                    cell.number_format = 'DD/MM/YYYY'
                    cell.alignment = Alignment(horizontal="center")
                else:
                    cell.alignment = Alignment(vertical="center")

        # Linha de total
        linha_total = len(df) + 3
        ws.cell(row=linha_total, column=2, value="Valor Total da Página:").font = Font(bold=True)

        for col_idx, campo in [(5, "valor_empenhado"), (6, "valor_liquidado"), (7, "valor_pago")]:
            total = df[campo].sum()
            cell = ws.cell(row=linha_total, column=col_idx, value=total)
            cell.number_format = '#,##0.00'
            cell.font = Font(bold=True, color=_AZ_ESCURO)
            cell.alignment = Alignment(horizontal="right")

        logger.info(
            "Aba '%s' atualizada: %d empenhos + linha de totais.", nome_aba, len(df)
        )

    # ------------------------------------------------------------------
    # Abas eFISCO (inserção / substituição)
    # ------------------------------------------------------------------

    def _injetar_aba_efisco(
        self,
        wb,
        nome: str,
        df: pd.DataFrame,
        cor_cabecalho: str,
    ) -> None:
        """
        Cria ou substitui a aba *nome* com os dados de *df*.
        Sempre inserida após as abas originais.
        """
        if nome in wb.sheetnames:
            del wb[nome]

        ws = wb.create_sheet(nome)
        ws.sheet_view.showGridLines = False

        if df.empty:
            ws["A1"] = f"(sem dados — coleta não executada ou banco indisponível)"
            ws["A1"].font = Font(italic=True, color="808080")
            return

        # Cabeçalho
        fill_cab  = PatternFill("solid", fgColor=cor_cabecalho)
        fonte_cab = Font(color=_BRANCO, bold=True, size=10)

        for col_idx, col in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col.upper().replace("_", " "))
            cell.fill = fill_cab
            cell.font = fonte_cab
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.row_dimensions[1].height = 28

        # Colunas monetárias e percentuais
        mon = {c for c in df.columns if any(
            x in c for x in ("valor", "dotacao", "saldo", "empenhado", "liquidado", "pago", "total")
        )}
        pct = {c for c in df.columns if "pct" in c}

        fill_par   = PatternFill("solid", fgColor="EBF3FB")
        fill_impar = PatternFill("solid", fgColor=_BRANCO)

        for row_idx, row_data in enumerate(df.itertuples(index=False, name=None), 2):
            fill = fill_par if row_idx % 2 == 0 else fill_impar
            for col_idx, valor in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=valor)
                cell.fill = fill
                col_nome = df.columns[col_idx - 1]
                if col_nome in mon:
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_nome in pct:
                    cell.number_format = '0.00"%"'
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(vertical="center")

        # Linha de totais
        lin_tot = len(df) + 2
        ws.cell(row=lin_tot, column=1, value="TOTAL").font = Font(bold=True, color=_BRANCO)
        ws.cell(row=lin_tot, column=1).fill = PatternFill("solid", fgColor=_AZ_ESCURO)

        for col_idx, col_nome in enumerate(df.columns, 1):
            cell_tot = ws.cell(row=lin_tot, column=col_idx)
            if col_nome in mon:
                col_letra = get_column_letter(col_idx)
                cell_tot.value = f"=SUM({col_letra}2:{col_letra}{lin_tot-1})"
                cell_tot.number_format = '#,##0.00'
                cell_tot.font = Font(bold=True, color=_BRANCO)
                cell_tot.alignment = Alignment(horizontal="right")
            cell_tot.fill = PatternFill("solid", fgColor=_AZ_ESCURO)

        # Ajusta largura
        for col_idx, col_nome in enumerate(df.columns, 1):
            max_len = max(
                len(col_nome),
                df[col_nome].astype(str).str.len().max() if not df.empty else 0,
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 50)

        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}1"
        ws.freeze_panes = "A2"

        logger.info("Aba '%s' gerada: %d linhas.", nome, len(df))

    # ------------------------------------------------------------------
    # Consolidado (view Power BI recriada localmente)
    # ------------------------------------------------------------------

    @staticmethod
    def _consolidado(df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            return pd.DataFrame()
        grp = df.groupby(
            ["unidade_gestora", "acao", "subacao", "data_coleta"], as_index=False
        ).agg(
            total_empenhado=("valor_empenhado", "sum"),
            total_liquidado=("valor_liquidado", "sum"),
            total_pago=("valor_pago", "sum"),
        )
        grp["saldo"] = grp["total_empenhado"] - grp["total_pago"]
        grp["pct_execucao"] = grp.apply(
            lambda r: round(r["total_pago"] / r["total_empenhado"] * 100, 2)
            if r["total_empenhado"] > 0 else 0.0,
            axis=1,
        )
        return grp.sort_values("total_empenhado", ascending=False).reset_index(drop=True)
